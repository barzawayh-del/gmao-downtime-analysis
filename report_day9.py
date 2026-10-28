import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

# 1. Load the clean data
df = pd.read_csv("gmao_clean.csv", parse_dates=["date"])

# 2. Compute the metrics (same logic as day 5)
days = (df["date"].max() - df["date"].min()).days + 1
period_h = days * 24                       # assumption: 24h/day operation
g = df.groupby("machine")["downtime_min"]
stats = pd.DataFrame({
    "failures": g.size(),
    "known_durations": g.count(),
    "downtime_h": (g.sum() / 60).round(2),
    "MTTR_min": g.mean().round(1),
})
stats["MTBF_h"] = ((period_h - stats["downtime_h"]) / stats["failures"]).round(1)
stats["MTTR_reliable"] = stats["known_durations"] >= 2
stats = stats.reset_index()

# 3. Write the Excel file with two sheets
with pd.ExcelWriter("gmao_report.xlsx", engine="openpyxl") as writer:
    stats.to_excel(writer, sheet_name="Summary", index=False)
    df.to_excel(writer, sheet_name="Clean log", index=False)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    for ws in writer.book.worksheets:
        for cell in ws[1]:                       # style the header row
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"                   # keep header visible when scrolling
        for col in ws.columns:                   # auto column width
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = width + 3

    # Dates without the time part
    for cell in writer.sheets["Clean log"]["A"][1:]:
        cell.number_format = "yyyy-mm-dd"

    # Bar chart of MTBF per machine
    ws = writer.sheets["Summary"]
    chart = BarChart()
    chart.title = "MTBF per machine (hours)"
    data = Reference(ws, min_col=6, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "I2")

print("Saved gmao_report.xlsx")
print(stats)
